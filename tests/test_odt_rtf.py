"""ODT und RTF liest und schreibt Lunivo selbst — ohne fremdes Programm im Hintergrund."""
import os
import subprocess
import sys
import time
import zipfile

import docx
import pytest
from PySide6.QtGui import QColor, QFont, QTextCharFormat, QTextCursor, QTextDocument, QTextListFormat

from lunivo_writer.models.page import HeaderFooter, PageSettings
from lunivo_writer.models.report import DocumentIOError
from lunivo_writer.services import file_service as formats


def _beispiel(tmp_path):
    d = docx.Document()
    d.add_heading("Überschrift", 1)
    p = d.add_paragraph("Normaler Text mit "); r = p.add_run("Fett"); r.bold = True
    d.add_paragraph("Punkt", style="List Bullet")
    t = d.add_table(rows=1, cols=2); t.cell(0, 0).text = "A"; t.cell(0, 1).text = "B"
    f = tmp_path / "q.docx"; d.save(f)
    return str(f)


@pytest.mark.parametrize("ext", [".odt", ".rtf"])
def test_roundtrip_ohne_umwandler(app, tmp_path, ext):
    l = formats.load(_beispiel(tmp_path))
    out = str(tmp_path / ("x" + ext))
    rep = formats.save(out, l.doc, l.page, l.meta)
    assert os.path.getsize(out) > 200 and not rep.warnings          # Speichern in ein anderes Format ohne Hinweisfenster
    back = formats.load(out)
    txt = back.doc.toPlainText()
    assert "Überschrift" in txt and "Normaler Text mit Fett" in txt and "Punkt" in txt and "A" in txt and "B" in txt
    assert not back.report.warnings
    assert back.doc.find("Fett").charFormat().fontWeight() >= 600
    assert back.doc.firstBlock().blockFormat().headingLevel() == 1
    punkt = back.doc.findBlock(back.doc.find("Punkt").position())
    assert punkt.textList() is not None
    assert back.doc.rootFrame().childFrames(), "Tabelle fehlt"


@pytest.mark.parametrize("ext", [".odt", ".rtf"])
def test_seite_kopf_und_fusszeile(app, tmp_path, ext):
    d = QTextDocument(); QTextCursor(d).insertText("Inhalt")
    page = PageSettings(width_mm=148, height_mm=210, margin_left_mm=20, margin_right_mm=15, margin_top_mm=30, margin_bottom_mm=28)
    page.header = HeaderFooter("Firma Muster", "right")
    page.footer = HeaderFooter("Seite {PAGE} von {PAGES}", "center")
    out = str(tmp_path / ("k" + ext))
    formats.save(out, d, page)
    p = formats.load(out).page
    assert abs(p.width_mm - 148) < 0.5 and abs(p.height_mm - 210) < 0.5
    assert abs(p.margin_left_mm - 20) < 0.5 and abs(p.margin_right_mm - 15) < 0.5
    assert abs(p.margin_top_mm - 30) < 1 and abs(p.margin_bottom_mm - 28) < 1
    assert p.header.text == "Firma Muster" and p.header.align == "right"
    assert p.footer.text == "Seite {PAGE} von {PAGES}" and p.footer.align == "center"


@pytest.mark.parametrize("ext", [".odt", ".rtf"])
def test_zeichenformat_umlaute_und_liste(app, tmp_path, ext):
    d = QTextDocument(); c = QTextCursor(d)
    f = QTextCharFormat(); f.setFontItalic(True); f.setForeground(QColor("#c00000")); f.setFontPointSize(16)
    c.insertText("Rot kursiv äöü ß € 𝄞", f)
    c.insertBlock(); c.insertText("Normal", QTextCharFormat())
    c.insertBlock(); lf = QTextListFormat(); lf.setStyle(QTextListFormat.ListDecimal); c.createList(lf); c.insertText("Erstens")
    c.insertBlock(); c.insertText("Zweitens")
    out = str(tmp_path / ("z" + ext))
    formats.save(out, d, PageSettings())
    back = formats.load(out).doc
    cf = back.find("Rot kursiv").charFormat()
    assert cf.fontItalic() and cf.foreground().color().name() == "#c00000" and round(cf.fontPointSize()) == 16
    assert "Rot kursiv äöü ß € 𝄞" in back.toPlainText()
    zweitens = back.findBlock(back.find("Zweitens").position())
    assert zweitens.textList() is not None and zweitens.textList().format().style() == QTextListFormat.ListDecimal
    assert zweitens.textList() is back.findBlock(back.find("Erstens").position()).textList()


def test_odt_paket_ist_gueltig(app, tmp_path):
    out = str(tmp_path / "p.odt")
    formats.save(out, QTextDocument("x"), PageSettings())
    with zipfile.ZipFile(out) as z:
        erster = z.infolist()[0]
        assert erster.filename == "mimetype" and erster.compress_type == zipfile.ZIP_STORED
        assert {"content.xml", "styles.xml", "META-INF/manifest.xml"} <= set(z.namelist())
        assert b"styles.xml" in z.read("META-INF/manifest.xml")


def test_fremdes_rtf_lesen(app, tmp_path):
    """Ein RTF, wie es Word schreibt: Codepage, \\u-Zeichen, Felder, Tabelle, Kopfzeile, Fußnote."""
    rtf = (rb"{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0\froman Times New Roman;}{\f1\fswiss Arial;}}"
           rb"{\colortbl;\red255\green0\blue0;}{\stylesheet{\s0 Normal;}{\s1 heading 1;}}"
           rb"{\info{\title Mein Titel}{\author Kay}}\paperw11906\paperh16838\margl1134\margr1134\margt1417\margb1134"
           rb"{\header\pard\qr Kopf\par}{\footer\pard\qc Seite {\field{\*\fldinst PAGE}{\fldrslt 1}}\par}"
           rb"\pard\s1\outlinelevel0\f1\fs32\b Titel\b0\par"
           rb"\pard\plain\f0\fs24 Gr\'fc\'dfe \u8364? und {\cf1 rot}{\footnote\pard Fu\'dfnote}.\par"
           rb"{\field{\*\fldinst HYPERLINK \"https://x\"}{\fldrslt Verweis}}\par"
           rb"\trowd\cellx3000\cellx6000\pard\intbl A1\cell B1\cell\row"
           rb"\pard Ende\par}")
    p = tmp_path / "word.rtf"; p.write_bytes(rtf)
    l = formats.load(str(p))
    t = l.doc.toPlainText()
    assert "Titel" in t and "Grüße € und rot." in t and "Verweis" in t and "A1" in t and "B1" in t and "Ende" in t
    assert "Fußnote" not in t and any("Endnoten" in w for w in l.report.warnings)
    assert l.doc.find("rot").charFormat().foreground().color().name() == "#ff0000"
    assert l.doc.firstBlock().blockFormat().headingLevel() == 1
    assert l.page.header.text == "Kopf" and l.page.header.align == "right" and l.page.footer.text == "Seite {PAGE}"
    assert abs(l.page.margin_left_mm - 20) < 0.1 and l.meta["title"] == "Mein Titel"


def test_doc_wird_klar_abgelehnt(app, tmp_path):
    p = tmp_path / "alt.doc"; p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600)
    with pytest.raises(DocumentIOError, match="nicht mehr unterstützt.*DOCX"):
        formats.load(str(p))
    assert ".doc" not in formats.SAVE_FILTERS


def test_dotx_vorlage_oeffnen(app, tmp_path):
    q = _beispiel(tmp_path)
    ziel = tmp_path / "vorlage.dotx"
    with zipfile.ZipFile(q) as z, zipfile.ZipFile(ziel, "w") as aus:
        for i in z.infolist():
            daten = z.read(i.filename)
            if i.filename == "[Content_Types].xml":
                daten = daten.replace(b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml")
            aus.writestr(i, daten)
    assert "Normaler Text mit Fett" in formats.load(str(ziel)).doc.toPlainText()


def test_kaputte_dateien(app, tmp_path):
    for name, inhalt in (("kaputt.odt", b"kein odt"), ("leer.odt", b"PK\x05\x06" + b"\0" * 18), ("kein.rtf", b"Hallo"),
                         ("kaputt.fodt", b"<?xml version='1.0'?><kaputt")):
        p = tmp_path / name; p.write_bytes(inhalt)
        with pytest.raises(DocumentIOError):
            formats.load(str(p))


def test_kein_fremdes_programm_beim_speichern_und_oeffnen(app, tmp_path, monkeypatch):
    """Weder beim Speichern noch beim Öffnen startet Lunivo einen anderen Prozess."""
    def verboten(*a, **k):
        raise AssertionError(f"Fremdes Programm gestartet: {a[:1]}")
    monkeypatch.setattr(subprocess, "Popen", verboten)
    monkeypatch.setattr(subprocess, "run", verboten)
    for ext in (".odt", ".rtf", ".docx", ".txt"):
        out = str(tmp_path / ("f" + ext))
        formats.save(out, QTextDocument("ohne Umweg"), PageSettings())
        assert "ohne Umweg" in formats.load(out).doc.toPlainText()


@pytest.mark.parametrize("ext", [".odt", ".rtf"])
def test_schnell_auch_bei_grossem_dokument(app, tmp_path, ext):
    d = QTextDocument(); c = QTextCursor(d)
    fett = QTextCharFormat(); fett.setFontWeight(QFont.Bold)
    for i in range(3000):
        c.insertText(f"Absatz {i} mit etwas Text, ", QTextCharFormat()); c.insertText("fett", fett); c.insertBlock()
    out = str(tmp_path / ("gross" + ext))
    t0 = time.monotonic()
    formats.save(out, d, PageSettings())
    gespeichert = time.monotonic() - t0
    back = formats.load(out)
    gesamt = time.monotonic() - t0
    assert "Absatz 2999" in back.doc.toPlainText()
    assert gespeichert < 5 and gesamt < 15, (gespeichert, gesamt)
    print(f"{ext}: speichern {gespeichert:.2f} s, mit Öffnen {gesamt:.2f} s", file=sys.stderr)


def test_odt_freie_bilder_kopf_fuss_und_text(app, tmp_path):
    """Bilder in Kopf- und Fußzeile und frei positionierte Bilder im Text gehen in ODT mit, ohne Hinweisfenster."""
    from PySide6.QtCore import QUrl
    from PySide6.QtGui import QImage
    from lunivo_writer.editor import freie_bilder as FB
    doc = QTextDocument(); cur = QTextCursor(doc); cur.insertText("Brief")
    rot = QImage(40, 20, QImage.Format_RGB32); rot.fill(QColor("red"))
    doc.addResource(QTextDocument.ImageResource, QUrl("frei-1"), rot)
    FB.an_absatz(cur, {"n": "frei-1", "b": 80, "h": 40, "bh": "seite", "bv": "seite", "xh": 100, "yv": 50, "ah": "", "av": "", "hinter": True})
    p = PageSettings(); p.header = HeaderFooter("Kopf", "center")
    logo = QImage(30, 30, QImage.Format_RGB32); logo.fill(QColor("blue"))
    p.bilder = [{"ort": "kopf", "bild": logo, "b": 60, "h": 60, "bh": "seite", "bv": "seite", "ah": "right", "av": "", "xh": 0, "yv": 20},
                {"ort": "fuss", "bild": logo, "b": 30, "h": 30, "bh": "rand", "bv": "seite", "ah": "center", "av": "bottom", "xh": 0, "yv": 0}]
    out = str(tmp_path / "bilder.odt")
    assert not formats.save(out, doc, p, {}).warnings
    back = formats.load(out)
    assert not back.report.warnings
    frei = FB.aus_absatz(back.doc.firstBlock().blockFormat())
    assert len(frei) == 1 and frei[0]["hinter"] and frei[0]["bh"] == "seite" and round(frei[0]["xh"]) == 100
    assert FB.bild_im_dokument(back.doc, frei[0]["n"]) is not None
    orte = sorted((a["ort"], a["ah"], a["av"]) for a in back.page.bilder)
    assert orte == [("fuss", "center", "bottom"), ("kopf", "right", "")]
    assert all(not a["bild"].isNull() for a in back.page.bilder)
